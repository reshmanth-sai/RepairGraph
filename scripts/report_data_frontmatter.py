"""
RepairGraph Report Data — Frontmatter & Preliminary Pages
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_frontmatter(doc, add_p, add_heading_1, add_heading_2, format_table):
    # =========================================================================
    # =========================================================================
    # TITLE PAGE — VIT CHENNAI CLOUD DA PROJECT
    # =========================================================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(28)
    p_pre.paragraph_format.space_after = Pt(6)
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_pre.add_run("VELLORE INSTITUTE OF TECHNOLOGY (VIT), CHENNAI\nSchool of Computer Science and Engineering\nCLOUD DIGITAL ASSIGNMENT (DA) PROJECT REPORT")
    r_inst.font.name = 'Arial'
    r_inst.font.size = Pt(11.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    p_course = doc.add_paragraph()
    p_course.paragraph_format.space_before = Pt(4)
    p_course.paragraph_format.space_after = Pt(16)
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_course = p_course.add_run("BACSE344: CLOUD INFRASTRUCTURE AND ARCHITECTURE (C1 SLOT)")
    r_course.font.name = 'Arial'
    r_course.font.size = Pt(12)
    r_course.font.bold = True
    r_course.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("REPAIRGRAPH")
    r_title.font.name = 'Arial'
    r_title.font.bold = True
    r_title.font.size = Pt(28)
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(20)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Cloud-Native Hardware Diagnostics, Repairability Assessment & Competitive Lifecycle Marketplace\nDeployed on Vercel Serverless Platform & Managed Neon PostgreSQL")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(12.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_before = Pt(6)
    p_desc.paragraph_format.space_after = Pt(28)
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_desc = p_desc.add_run("A Cloud Computing Project Report submitted for the Digital Assignment (DA) Evaluation in\nBACSE344: Cloud Infrastructure and Architecture\nVIT Chennai")
    r_desc.font.name = 'Arial'
    r_desc.font.size = Pt(10.5)
    r_desc.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Student & Course Evaluation Table
    tbl = doc.add_table(rows=5, cols=2)
    format_table(tbl, col_widths=[3.1, 3.1])
    
    rows_data = [
        ("Submitted By:", "Course & Faculty Details:"),
        ("NAIDU RESHMANTH SAI", "Faculty: Dr. P. Anandan"),
        ("Register No: 25BCE1112", "Course: Cloud Infrastructure and Architecture"),
        ("Slot: C1 Slot | B.Tech CSE", "Course Code: BACSE344 (C1 Slot)"),
        ("Institution: VIT Chennai", "Submission Date: 12-09-2026\nProduction: https://repairgraph.vercel.app")
    ]
    for i, (col1, col2) in enumerate(rows_data):
        cell1 = tbl.cell(i, 0)
        cell2 = tbl.cell(i, 1)
        cell1.paragraphs[0].text = col1
        cell2.paragraphs[0].text = col2
        if i == 0:
            cell1.paragraphs[0].runs[0].font.bold = True
            cell2.paragraphs[0].runs[0].font.bold = True
        elif i == 1:
            cell1.paragraphs[0].runs[0].font.bold = True
            cell1.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_page_break()

    # =========================================================================
    # ABSTRACT
    # =========================================================================
    add_heading_1(doc, "ABSTRACT")
    add_p(doc, 
          "The modern consumer electronics ecosystem suffers from structural market inefficiencies, opaque repair pricing, "
          "artificial hardware locking, and premature device retirement that directly drives escalating global electronic "
          "waste. While statutory initiatives such as India's Right to Repair Portal (Department of Consumer Affairs) and the "
          "E-Waste (Management) Rules, 2022 mandate fair access to repair and extended producer responsibility, consumers "
          "and enterprise fleet managers lack an objective, evidence-based platform to triage hardware failures and evaluate "
          "the economic viability of repair versus replacement.")
    
    add_p(doc, 
          "This report presents RepairGraph, a production-grade, cloud-native hardware diagnostics, repairability evaluation, "
          "competitive quote marketplace, and standardized device service passport platform. Built as a unified full-stack "
          "Next.js 15 application deployed on Vercel Serverless Runtime (Node.js 24.x) and backed by managed Neon Serverless "
          "PostgreSQL via Prisma ORM, RepairGraph integrates a two-tier hybrid intelligence architecture.")
    
    add_p(doc, 
          "Rather than relying on unconstrained, hallucination-prone large language models for financial and statutory decisions, "
          "RepairGraph implements a fully deterministic rule engine. Tier 1 performs heuristic word-boundary symptom parsing "
          "and hazard classification across eight hardware subsystems. Tier 2 evaluates an original 7-Factor Repairability Score "
          "Model (Disassembly, Parts Availability, Documentation, Modularity, Software Pairing, Age/Lifecycle, and Local Service "
          "Ecosystem summing to 100 points), brand-scaled component replacement matrices, and a Repair Cost Ratio (RCR) model. "
          "The engine outputs actionable verdicts: REPAIR, DIY, RESELL, REPLACE, or RECYCLE, accompanied by transparent, audit-ready "
          "reasoning citing manufacturing carbon offsets and statutory e-waste divert estimates.")
    
    add_p(doc, 
          "Furthermore, RepairGraph provides a verified two-sided marketplace where verified independent repair workshops submit "
          "competitive bids against consumer tickets. Upon job acceptance and completion, the platform immutably records "
          "standardized maintenance milestones into a physical device Repair Passport tied to hardware serial numbers.")
    
    add_p(doc, 
          "The system's engineering rigor is substantiated by 82 automated test cases (35 deterministic engine unit tests + "
          "47 business rule and security integration tests, 100% passing), clean ESLint validation, strict zero-leak error "
          "sanitization, and live end-to-end acceptance validation on its production deployment at "
          "https://repairgraph.vercel.app.")
    
    add_p(doc, 
          "Keywords: Hardware Diagnostics, Right to Repair, E-Waste Management, Deterministic Decision Engine, "
          "Repairability Scoring, Cloud Computing, Next.js, Vercel Serverless, Neon PostgreSQL, Prisma ORM.", 
          bold_prefix="Keywords: ", space_after=18)

    doc.add_page_break()

    # =========================================================================
    # TABLE OF CONTENTS
    # =========================================================================
    add_heading_1(doc, "TABLE OF CONTENTS")
    
    toc_data = [
        ("Abstract", "i"),
        ("List of Figures", "iii"),
        ("List of Tables", "iv"),
        ("Chapter 1: Introduction", "1"),
        ("    1.1 Background & Industry Landscape", "1"),
        ("    1.2 Problem Context: The Right to Repair & E-Waste Crisis", "2"),
        ("    1.3 Motivation", "3"),
        ("    1.4 Problem Statement", "4"),
        ("    1.5 Objectives of the Platform", "4"),
        ("    1.6 Scope of the Project", "5"),
        ("    1.7 Key Engineering Contributions", "6"),
        ("Chapter 2: Existing System and Proposed System", "7"),
        ("    2.1 Existing Repair Approaches", "7"),
        ("    2.2 Limitations of the Current Landscape", "8"),
        ("    2.3 Proposed RepairGraph System", "9"),
        ("    2.4 Distinct Advantages of RepairGraph", "10"),
        ("    2.5 Comparative Analysis Matrix", "11"),
        ("Chapter 3: Requirements Analysis", "12"),
        ("    3.1 Functional Requirements", "12"),
        ("    3.2 Non-Functional Requirements", "13"),
        ("    3.3 Hardware Requirements", "14"),
        ("    3.4 Software Requirements", "15"),
        ("    3.5 User Roles & Permissions Matrix", "16"),
        ("    3.6 System Constraints & Invariants", "17"),
        ("Chapter 4: System Architecture", "18"),
        ("    4.1 Overall Architectural Topology", "18"),
        ("    4.2 Client Layer (React 19 & Tailwind CSS v4)", "19"),
        ("    4.3 Application & Service Layer", "20"),
        ("    4.4 API & Routing Layer", "21"),
        ("    4.5 Authentication & Session Management", "22"),
        ("    4.6 Deterministic Decision Engine", "22"),
        ("    4.7 Persistence Layer (Prisma ORM)", "23"),
        ("    4.8 Cloud Deployment Architecture", "24"),
        ("    4.9 End-to-End Data & Control Flow", "25"),
        ("Chapter 5: Technology Stack", "26"),
        ("    5.1 Framework & Runtime Selection", "26"),
        ("    5.2 Frontend & UI Libraries", "27"),
        ("    5.3 Backend & Validation Technologies", "27"),
        ("    5.4 Database & Persistence Stack", "28"),
        ("    5.5 Cloud Infrastructure Providers", "29"),
        ("    5.6 Technology Stack Inventory Table", "30"),
        ("Chapter 6: Functional Modules", "31"),
        ("    6.1 User Authentication & Session Control", "31"),
        ("    6.2 Hardware Device Catalog & Valuation", "32"),
        ("    6.3 Problem Reporting & Triage Initiation", "33"),
        ("    6.4 Heuristic Symptom Extraction", "34"),
        ("    6.5 Diagnostic Signal Synthesis", "34"),
        ("    6.6 7-Factor Repairability Evaluation", "35"),
        ("    6.7 Fair-Market Economic Analysis", "36"),
        ("    6.8 Statutory Decision & Explanation Generation", "37"),
        ("    6.9 Verified Repairer Directory & Geolocation", "38"),
        ("    6.10 Competitive Quote Marketplace", "39"),
        ("    6.11 Quote Lifecycle & Acceptance State Machine", "40"),
        ("    6.12 Repair Job Execution Stepper", "41"),
        ("    6.13 Standardized Repair Passport Ledger", "42"),
        ("    6.14 Post-Service Reviews & Rating Synthesis", "43"),
        ("    6.15 Administrative & Seed Controls", "44"),
        ("Chapter 7: RepairGraph Decision Engine", "45"),
        ("    7.1 Deterministic Architecture vs. Black-Box AI", "45"),
        ("    7.2 Token-Boundary Symptom Parsing", "46"),
        ("    7.3 Subsystem Classification & Hazard Detection", "47"),
        ("    7.4 Diagnostic Confidence Formulation", "48"),
        ("    7.5 7-Factor Baseline Repairability Score Model", "49"),
        ("    7.6 Incident-Specific Deduction Penalties", "51"),
        ("    7.7 Device Depreciation & Valuation Formulation", "52"),
        ("    7.8 Component Replacement Cost Benchmarks", "53"),
        ("    7.9 Repair Cost Ratio (RCR) Calculation", "54"),
        ("    7.10 Value-Preservation Economic Score", "54"),
        ("    7.11 Statutory Lifecycle Decision Matrix", "55"),
        ("    7.12 Automated Explainability Audit Trail", "56"),
        ("Chapter 8: Database Design", "57"),
        ("    8.1 Relational Design Philosophy", "57"),
        ("    8.2 Entity Relationship Diagram (11 Models)", "58"),
        ("    8.3 Relational Entity Specifications", "59"),
        ("    8.4 Custom Enums (9 Enums)", "61"),
        ("    8.5 Relational Invariants & Integrity Constraints", "62"),
        ("    8.6 Prisma Schema & Migration Baseline", "63"),
        ("    8.7 Cloud Serverless Pooling with Neon", "64"),
        ("Chapter 9: REST API Design", "65"),
        ("    9.1 RESTful Conventions & Error Payloads", "65"),
        ("    9.2 API Resource Namespaces", "66"),
        ("    9.3 HTTP Method Operations Catalog", "67"),
        ("    9.4 Input Validation with Zod", "69"),
        ("    9.5 IDOR Protection & Ownership Verification", "70"),
        ("Chapter 10: Authentication and Security", "71"),
        ("    10.1 Password Cryptography (bcrypt)", "71"),
        ("    10.2 JWT Session Tokens & Cookie Security", "71"),
        ("    10.3 Token Suppression Strategy", "72"),
        ("    10.4 In-Memory Sliding-Window Rate Limiting", "73"),
        ("    10.5 HTTP Defense-in-Depth Security Headers", "74"),
        ("    10.6 Error Sanitization & Leak Suppression", "75"),
        ("    10.7 Zero-Secret Version Control Policy", "75"),
        ("Chapter 11: User Workflow", "76"),
        ("    11.1 Primary Consumer Repair Journey", "76"),
        ("    11.2 Technician Quoting & Workbench Stepper", "77"),
        ("    11.3 Alternative Lifecycle Paths (DIY, Resell, Replace, Recycle)", "78"),
        ("Chapter 12: Implementation Architecture", "79"),
        ("    12.1 Project Directory Organization", "79"),
        ("    12.2 Layered Backend Architecture", "80"),
        ("    12.3 Database Singleton Pattern", "81"),
        ("Chapter 13: Testing and Validation", "82"),
        ("    13.1 Comprehensive Testing Strategy", "82"),
        ("    13.2 Unit Testing: Diagnostic & Decision Engine (35/35)", "83"),
        ("    13.3 Integration Testing: Backend & Security (47/47)", "84"),
        ("    13.4 Total Automated Test Audit (82/82 Passing)", "85"),
        ("    13.5 Production End-to-End Acceptance Validation", "86"),
        ("Chapter 14: Cloud Deployment", "87"),
        ("    14.1 Production Cloud Architecture", "87"),
        ("    14.2 Vercel Serverless Runtime (Node.js 24.x)", "88"),
        ("    14.3 Managed Neon Serverless PostgreSQL", "89"),
        ("    14.4 Environment Configuration & Secrets", "90"),
        ("    14.5 Production Deployment Runbook", "91"),
        ("    14.6 Health Telemetry Endpoint (/api/health)", "92"),
        ("    14.7 Free-Tier Cost Analysis", "93"),
        ("    14.8 Live Vercel Production Deployment Evidence", "94"),
        ("    14.9 Live Neon Serverless PostgreSQL Database Verification", "95"),
        ("    14.10 Live Platform User Interface Gallery & Operational Walkthrough", "96"),
        ("        14.10.1 Overview & System Telemetry Dashboard", "96"),
        ("        14.10.2 Identity Authentication & Multi-Persona Evaluator Switcher", "97"),
        ("        14.10.3 Competitive Repair Marketplace & Ticket Lifecycle Tracking", "97"),
        ("        14.10.4 Deterministic Diagnostic Intake & 7-Factor Scoring Suite", "98"),
        ("        14.10.5 Registered Hardware Inventory & Lifecycle Catalog", "98"),
        ("        14.10.6 Digital Product Passport & Standardized Maintenance Ledger", "99"),
        ("        14.10.7 Verified Technician Workbench & Sequential Milestone Stepper", "99"),
        ("Chapter 15: Results and Discussion", "100"),
        ("    15.1 Qualitative Diagnostic Evaluation", "100"),
        ("    15.2 Marketplace Efficiency Observations", "101"),
        ("    15.3 Regulatory Alignment & E-Waste Impact", "102"),
        ("Chapter 16: Limitations", "103"),
        ("    16.1 In-Memory Rate Limiting", "103"),
        ("    16.2 Academic Demonstration Workload Scope", "103"),
        ("    16.3 Synchronous Decision Evaluation", "104"),
        ("    16.4 Provider-Native Telemetry Reliance", "104"),
        ("    16.5 Non-Destructive Production Testing Boundary", "104"),
        ("    16.6 Informational Specialist Context", "105"),
        ("Chapter 17: Future Enhancements", "106"),
        ("    17.1 Distributed Rate Limiting (Redis)", "106"),
        ("    17.2 Asynchronous Background Queue (BullMQ)", "106"),
        ("    17.3 Enterprise APM & Distributed Tracing", "107"),
        ("    17.4 Vector Embeddings for OEM Schematics", "107"),
        ("    17.5 Dedicated Mobile Application", "107"),
        ("Chapter 18: Conclusion", "108"),
        ("    18.1 Summary of Engineering Achievements", "108"),
        ("    18.2 Practical Impact & Academic Defense Summary", "109"),
        ("References", "110"),
        ("Appendices", "112"),
        ("    Appendix A: Complete REST API Endpoint Inventory", "112"),
        ("    Appendix B: Relational Data Model Specifications", "115"),
        ("    Appendix C: Diagnostic Engine Formulas & Matrix Weights", "118"),
        ("    Appendix D: Automated Test Suite Inventory (82 Tests)", "120"),
        ("    Appendix E: Production Deployment Configuration Summary", "124")
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_data), cols=2)
    format_table(tbl_toc, col_widths=[5.2, 1.0])
    for idx, (title, page) in enumerate(toc_data):
        c1 = tbl_toc.cell(idx, 0)
        c2 = tbl_toc.cell(idx, 1)
        c1.paragraphs[0].text = title
        c2.paragraphs[0].text = page
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if title.startswith("Chapter") or title in ["References", "Appendices"]:
            c1.paragraphs[0].runs[0].font.bold = True
            c2.paragraphs[0].runs[0].font.bold = True

    doc.add_page_break()

    # =========================================================================
    # LIST OF FIGURES & LIST OF TABLES
    # =========================================================================
    add_heading_1(doc, "LIST OF FIGURES")
    figures = [
        ("Figure 1", "RepairGraph Unified Full-Stack Architecture", "18"),
        ("Figure 2", "End-to-End Request & Data Processing Pipeline", "25"),
        ("Figure 3", "Deterministic Diagnostic & Decision Engine Flow", "56"),
        ("Figure 4", "Relational Database Schema & Entity Relationships (11 Models)", "58"),
        ("Figure 5", "Competitive Repair Marketplace & Lifecycle Workflow", "76"),
        ("Figure 6", "Production Cloud Deployment & Network Security Topology", "87"),
        ("Figure 7", "Vercel Cloud Production Deployment Dashboard & Build Telemetry (Deployment GXaQwF4dR)", "94"),
        ("Figure 8", "Neon Cloud PostgreSQL Console & Normalized Relational Tables (Public Schema)", "95"),
        ("Figure 9", "RepairGraph Production Overview & Telemetry Dashboard (https://repairgraph.vercel.app)", "96"),
        ("Figure 10", "Identity Authentication & Multi-Persona Evaluator Switcher (/login)", "97"),
        ("Figure 11", "Competitive Repair Marketplace & Ticket Lifecycle Tracking (/repairs)", "97"),
        ("Figure 12", "Deterministic Diagnostic Intake & 7-Factor Repairability Evaluation (/report)", "98"),
        ("Figure 13", "Hardware Inventory Catalog & Lifecycle Valuation Management (/devices)", "98"),
        ("Figure 14", "Digital Product Passport & Standardized Maintenance Ledger (/passport)", "99"),
        ("Figure 15", "Verified Technician Workbench & Sequential Milestone Stepper (/repairer)", "99")
    ]
    tbl_fig = doc.add_table(rows=len(figures), cols=3)
    format_table(tbl_fig, col_widths=[1.2, 4.2, 0.8])
    for idx, (num, title, pg) in enumerate(figures):
        tbl_fig.cell(idx, 0).paragraphs[0].text = num
        tbl_fig.cell(idx, 1).paragraphs[0].text = title
        tbl_fig.cell(idx, 2).paragraphs[0].text = pg
        tbl_fig.cell(idx, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    add_heading_1(doc, "LIST OF TABLES")
    tables = [
        ("Table 1", "Comparative Matrix: Traditional Repair vs. RepairGraph", "11"),
        ("Table 2", "Hardware & Software Execution Requirements", "15"),
        ("Table 3", "User Roles and Permission Boundaries", "16"),
        ("Table 4", "Technology Stack and Operational Roles", "30"),
        ("Table 5", "7-Factor Design Repairability Weight Distribution", "50"),
        ("Table 6", "Incident-Specific Penalty Deduction Rules", "51"),
        ("Table 7", "Statutory Lifecycle Action Decision Rules", "55"),
        ("Table 8", "11 Normalized Relational Entities in Prisma Schema", "59"),
        ("Table 9", "Prisma Database Enums Specification", "61"),
        ("Table 10", "CRUD Functionality Across Resource Namespaces", "66"),
        ("Table 11", "Representative REST API Route Catalog", "67"),
        ("Table 12", "Automated Test Suite Execution Summary (82 Tests)", "85"),
        ("Table 13", "Production End-to-End Acceptance Verification Summary", "86"),
        ("Table 14", "Environment Variables Specification (Names Only)", "90")
    ]
    tbl_tbl = doc.add_table(rows=len(tables), cols=3)
    format_table(tbl_tbl, col_widths=[1.2, 4.2, 0.8])
    for idx, (num, title, pg) in enumerate(tables):
        tbl_tbl.cell(idx, 0).paragraphs[0].text = num
        tbl_tbl.cell(idx, 1).paragraphs[0].text = title
        tbl_tbl.cell(idx, 2).paragraphs[0].text = pg
        tbl_tbl.cell(idx, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_page_break()
