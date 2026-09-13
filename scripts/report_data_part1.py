"""
RepairGraph Report Data — Part 1: Chapters 1 through 5
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_part1(doc, add_p, add_bullet, add_heading_1, add_heading_2, add_heading_3, 
                add_figure_image, add_callout_box, format_table):
    
    # =========================================================================
    # CHAPTER 1 — INTRODUCTION
    # =========================================================================
    add_heading_1(doc, "CHAPTER 1: INTRODUCTION")
    
    add_heading_2(doc, "1.1 Background & Industry Landscape")
    add_p(doc, 
          "The consumer electronics and computing hardware sectors have witnessed unprecedented expansion over the past two "
          "decades, driven by rapid miniaturization, high-density system-on-chip (SoC) integration, and aggressive annual product "
          "release cycles. Modern laptops, smartphones, tablets, and wearable devices pack computational capabilities that rival "
          "enterprise workstations of earlier eras into ultra-thin enclosures. However, this engineering trajectory has been "
          "accompanied by an aggressive shift away from repair-friendly industrial designs toward adhesive-bonded assemblies, "
          "soldered unified memory architectures, proprietary tamper-resistant fasteners, and cryptographic component serialization.")
    add_p(doc, 
          "When a consumer hardware unit experiences failure—whether a cracked display digitizer, chemically degraded lithium-ion cell, "
          "or thermal cooling fan bearing seizure—the owner encounters an opaque, highly friction-laden repair journey. Authorized "
          "Original Equipment Manufacturer (OEM) service centers frequently quote exorbitant repair estimates that approach 70% to 90% "
          "of the device's original retail purchase price. This pricing strategy systematically steers consumers toward premature device "
          "abandonment and new device acquisition, contributing directly to the worsening global electronic waste crisis.")

    add_heading_2(doc, "1.2 Problem Context: The Right to Repair & E-Waste Crisis")
    add_p(doc, 
          "The statutory and environmental backdrop of device servicing has reached a critical turning point globally and within India. "
          "Under the auspices of the Department of Consumer Affairs, Government of India, the national Right to Repair Portal was "
          "launched to curb planned obsolescence, mandate OEM transparency regarding diagnostic manuals and spare parts, and foster "
          "a competitive ecosystem of certified third-party technicians. Concurrently, India's statutory E-Waste (Management) Rules, 2022 "
          "impose strict Extended Producer Responsibility (EPR) mandates, requiring transparent lifecycle tracking, safe recovery of hazardous "
          "substances (such as lead, cadmium, and brominated flame retardants), and diversion of non-repairable electronic hardware into "
          "authorized recycling streams.")
    add_p(doc, 
          "Despite these statutory strides, consumers and enterprise fleet administrators lack a standardized software infrastructure "
          "capable of objectively evaluating hardware repairability, calculating real-world economic feasibility, comparing competitive "
          "local repair bids, and preserving permanent service provenance.")

    add_heading_2(doc, "1.3 Motivation")
    add_p(doc, 
          "The core motivation behind RepairGraph stems from three pervasive structural problems in consumer electronics ownership:")
    add_bullet(doc, "Asymmetric Information & Quoting Guesswork: Consumers rarely possess the technical diagnostic knowledge to verify whether an OEM's quote of ₹18,000 for a motherboard replacement is legitimate, or whether a ₹3,500 repair at an independent workshop is technically viable and economically sound.", bold_prefix="1. ")
    add_bullet(doc, "Unconstrained AI Hallucinations in Consumer Tech: Generic generative AI chatbots frequently fabricate repair procedures, misquote non-existent component prices, or hallucinate dangerous advice regarding swollen lithium-ion cells, creating physical safety hazards.", bold_prefix="2. ")
    add_bullet(doc, "Absence of Device Maintenance Provenance: Secondary electronics markets suffer severe value depreciation because prospective buyers cannot verify whether a refurbished or used laptop was professionally serviced with OEM-grade parts or compromised with substandard soldering.", bold_prefix="3. ")

    add_heading_2(doc, "1.4 Problem Statement")
    add_p(doc, 
          "To design, implement, test, and deploy a production-grade, cloud-native software platform that provides deterministic, "
          "rule-based hardware diagnostic triage, an objective 7-factor repairability scoring model, fair-market economic repair-vs-replace "
          "decision support, a verified two-sided competitive repair marketplace, and an immutable device service passport ledger "
          "accessible through an authenticated, secure web application.")

    add_heading_2(doc, "1.5 Objectives of the Platform")
    add_bullet(doc, "Develop a deterministic diagnostic parser that extracts failure modes, severity levels, and acute safety hazards (liquid ingress, battery swelling) from user-submitted symptom narratives without relying on probabilistic generative LLMs.", bold_prefix="• ")
    add_bullet(doc, "Formulate and calibrate an original 7-Factor Repairability Score Model (Disassembly, Parts Availability, Documentation, Modularity, Software Pairing, Age/Lifecycle, Service Ecosystem) summing to 100 points.", bold_prefix="• ")
    add_bullet(doc, "Engineer an economic valuation engine that computes consumer hardware depreciation (22% annual curve), derives Repair Cost Ratios (RCR), and maps outcomes to five statutory lifecycle recommendations: REPAIR, DIY, RESELL, REPLACE, or RECYCLE.", bold_prefix="• ")
    add_bullet(doc, "Construct a secure, relational marketplace enabling verified independent repair workshops to submit competitive quotes and execute structured repair jobs through a verified state machine.", bold_prefix="• ")
    add_bullet(doc, "Provide a standardized, tamper-evident Repair Passport linked to physical device serial numbers that permanently records completed service history and parts replacements.", bold_prefix="• ")
    add_bullet(doc, "Deploy the platform to a modern cloud-computing infrastructure utilizing Vercel Serverless Runtime (Node.js 24.x) and Neon Managed Serverless PostgreSQL with automated connection pooling and end-to-end TLS encryption.", bold_prefix="• ")

    add_heading_2(doc, "1.6 Scope of the Project")
    add_p(doc, 
          "The scope encompasses five major consumer hardware categories widely prevalent in academic and enterprise environments: "
          "Laptops, Smartphones, Tablets, Headphones, and Monitors. The platform provides full lifecycle support from initial user onboarding "
          "and device registration, through interactive symptom triage and quote acceptance, to post-service verification and customer reviews. "
          "The scope explicitly excludes automated hardware disassembly robots, direct online payment gateway settlement, and speculative "
          "blockchain mechanisms, focusing strictly on auditable software engineering and statutory compliance.")

    add_heading_2(doc, "1.7 Key Engineering Contributions")
    add_bullet(doc, "Two-Tier Hybrid Architecture: Separates deterministic heuristic symptom parsing (Tier 1) from mathematical scoring and statutory decision trees (Tier 2), ensuring 100% reproducible diagnostic outputs with sub-millisecond execution latency.", bold_prefix="1. ")
    add_bullet(doc, "Statutory Alignment with Indian Environmental Law: Integrates statutory mandates from India's E-Waste (Management) Rules, 2022 and Right to Repair Portal directly into automated decision matrices and explainability narratives.", bold_prefix="2. ")
    add_bullet(doc, "Production Cloud-Native Unified Architecture: Eliminates multi-service synchronization overhead by combining Next.js 15 App Router frontend and REST API route handlers into a single Vercel serverless deployment connected to Neon PostgreSQL.", bold_prefix="3. ")
    add_bullet(doc, "Defense-in-Depth Security Perimeter: Incorporates HttpOnly/SameSite session cookies, sliding-window rate limiting, Zod schema validation, IDOR authorization barriers, and zero-leak database error sanitization.", bold_prefix="4. ")
    add_bullet(doc, "Exhaustive Validation Rigor: 82 automated test cases (35 deterministic engine unit tests + 47 backend and security integration tests, 0 failures) verified in local and live production environments.", bold_prefix="5. ")

    # =========================================================================
    # CHAPTER 2 — EXISTING SYSTEM AND PROPOSED SYSTEM
    # =========================================================================
    add_heading_1(doc, "CHAPTER 2: EXISTING SYSTEM AND PROPOSED SYSTEM")

    add_heading_2(doc, "2.1 Existing Repair Approaches")
    add_p(doc, 
          "The current landscape for consumer electronics repair consists of three fragmented channels:")
    add_bullet(doc, "Authorized OEM Service Depots: High repair charges, proprietary parts gating, lengthy turnaround times (often 7–14 days), and aggressive commercial steering toward trade-ins or replacement.", bold_prefix="1. ")
    add_bullet(doc, "Informal Street Workshops: Dense physical presence and lower prices, but suffering from zero parts provenance transparency, unverified technician credentials, absence of written warranties, and no standardized maintenance records.", bold_prefix="2. ")
    add_bullet(doc, "Generic AI & Search Engines: Generalist search results and generative LLMs that provide non-deterministic troubleshooting steps, frequently confusing distinct hardware revisions, failing to flag swollen battery combustion risks, and lacking pricing context.", bold_prefix="3. ")

    add_heading_2(doc, "2.2 Limitations of the Current Landscape")
    add_bullet(doc, "Lack of Objective Economic Feasibility Metrics: Consumers have no mathematical framework to evaluate whether a ₹7,000 repair quote makes sense for a 3.5-year-old laptop worth ₹16,000 on the secondary market.", bold_prefix="• ")
    add_bullet(doc, "Absence of Competitive Quote Discovery: Getting comparative estimates requires physically transporting non-functional hardware across multiple disparate workshops across metropolitan hubs.", bold_prefix="• ")
    add_bullet(doc, "Fragile Maintenance Records: Paper invoices and thermal receipts fade or are lost within months, destroying secondary resale value and hampering warranty claims.", bold_prefix="• ")
    add_bullet(doc, "Environmental Disconnect: No existing consumer tool computes or visualizes the electronic waste diversion or carbon emissions prevented by choosing repair over disposal.", bold_prefix="• ")

    add_heading_2(doc, "2.3 Proposed System: RepairGraph")
    add_p(doc, 
          "RepairGraph addresses these structural shortcomings by introducing an integrated, cloud-native platform that unifies "
          "hardware asset management, deterministic diagnostic intelligence, competitive marketplace bidding, and permanent service "
          "passports. Users catalog their hardware assets once, receive instant deterministic repairability scores and fair-market economic "
          "analyses, solicit competitive quotes from verified independent technicians, track physical bench repair milestones, and receive "
          "a digitally signed Repair Passport upon completion.")

    add_heading_2(doc, "2.4 Distinct Advantages of RepairGraph")
    add_bullet(doc, "Deterministic Reliability: 100% reproducible scoring without generative model hallucinations.", bold_prefix="• ")
    add_bullet(doc, "Statutory Grounding: Directly cites and adheres to India's Right to Repair framework and E-Waste Rules 2022.", bold_prefix="• ")
    add_bullet(doc, "Consumer Protection: In-memory rate limiting, HttpOnly authentication cookies, and strict IDOR prevention.", bold_prefix="• ")
    add_bullet(doc, "Two-Sided Utility: Empowers both consumers seeking transparent service and independent technicians expanding their business.", bold_prefix="• ")

    add_heading_2(doc, "2.5 Comparative Analysis Matrix")
    add_p(doc, "Table 1 compares the traditional electronics repair experience against the capabilities provided by RepairGraph.")

    tbl_comp = doc.add_table(rows=8, cols=3)
    format_table(tbl_comp, col_widths=[1.8, 2.2, 2.2])
    comp_data = [
        ("Capability / Dimension", "Traditional Repair Ecosystem", "RepairGraph Platform"),
        ("Diagnostic Triage", "Subjective technician inspection or search", "Deterministic 2-Tier Rule Engine"),
        ("Repairability Evaluation", "Opaque OEM scoring or subjective opinion", "Transparent 7-Factor Model (0–100 Pts)"),
        ("Economic Valuation", "None (Arbitrary repair pricing)", "Automated 22%/yr Depreciation & RCR Math"),
        ("Quote Discovery", "Manual physical workshop visits", "Digital competitive bidding marketplace"),
        ("Service History Tracking", "Fading physical paper slips", "Immutable digital Repair Passport by serial"),
        ("E-Waste Regulatory Link", "Completely unaddressed", "Automated E-Waste Rules 2022 compliance"),
        ("Security Perimeter", "Unsecured phone/messaging communication", "JWT, HttpOnly cookies, Zod validation, HSTS")
    ]
    for idx, (c1, c2, c3) in enumerate(comp_data):
        tbl_comp.cell(idx, 0).paragraphs[0].text = c1
        tbl_comp.cell(idx, 1).paragraphs[0].text = c2
        tbl_comp.cell(idx, 2).paragraphs[0].text = c3

    # =========================================================================
    # CHAPTER 3 — REQUIREMENTS ANALYSIS
    # =========================================================================
    add_heading_1(doc, "CHAPTER 3: REQUIREMENTS ANALYSIS")

    add_heading_2(doc, "3.1 Functional Requirements")
    add_bullet(doc, "FR-1 User Authentication: Secure user registration and login with bcrypt password hashing (10 rounds) and JWT generation stored in HttpOnly cookies.", bold_prefix="• ")
    add_bullet(doc, "FR-2 Hardware Catalog Management: Full CRUD operations for registered devices, tracking serial numbers, purchase price, purchase date, condition, and fair market value.", bold_prefix="• ")
    add_bullet(doc, "FR-3 Diagnostic Problem Reporting: Ticket filing interface allowing users to input free-text symptom narratives, urgency levels, and observed failure modes.", bold_prefix="• ")
    add_bullet(doc, "FR-4 Automated Diagnostic & Economic Evaluation: Deterministic calculation of 7-factor repairability scores, incident penalties, depreciated market valuation, and statutory lifecycle actions.", bold_prefix="• ")
    add_bullet(doc, "FR-5 Technician Directory & Geolocation: Browsable listing of verified independent repair workshops with granular brand and component specializations.", bold_prefix="• ")
    add_bullet(doc, "FR-6 Competitive Quote Bidding: Capability for verified repairers to inspect open repair tickets and submit estimated costs, turnaround days, and scope notes.", bold_prefix="• ")
    add_bullet(doc, "FR-7 Quote Decision State Machine: Customer workflow to accept or reject quotes; quote acceptance automatically creates an active RepairJob and marks competing quotes rejected.", bold_prefix="• ")
    add_bullet(doc, "FR-8 Workbench Job Stepper: Multi-stage bench tracking interface (DIAGNOSING -> WAITING_FOR_PART -> REPAIRING -> TESTING -> COMPLETED) restricted strictly to assigned specialists.", bold_prefix="• ")
    add_bullet(doc, "FR-9 Permanent Repair Passport: Automated creation of immutable service records linked to physical hardware serial numbers upon job completion.", bold_prefix="• ")
    add_bullet(doc, "FR-10 Verified Reviews: Post-service rating (1–5 stars) and review submission restricted strictly to customers who have completed a repair job.", bold_prefix="• ")

    add_heading_2(doc, "3.2 Non-Functional Requirements")
    add_bullet(doc, "Sub-Millisecond Engine Latency: Deterministic diagnostic evaluations must execute in < 1 ms without external network hops or API dependencies.", bold_prefix="• Performance: ")
    add_bullet(doc, "Comprehensive Perimeter Defense: Strict suppression of raw JWTs from browser JSON bodies, sliding-window rate limiting on auth routes, and automated database error sanitization.", bold_prefix="• Security: ")
    add_bullet(doc, "Transactional Integrity: Relational foreign key constraints, cascading deletions, and unique invariants enforced by PostgreSQL to eliminate orphaned records.", bold_prefix="• Reliability: ")
    add_bullet(doc, "Zero-Downtime Serverless Scaling: Full compatibility with serverless auto-scaling and connection pooling (PgBouncer) on Neon PostgreSQL.", bold_prefix="• Scalability: ")
    add_bullet(doc, "Responsive Editorial User Interface: Mobile-first responsive layouts styled with Tailwind CSS v4, supporting reduced-motion and high-contrast accessibility standards.", bold_prefix="• Usability: ")

    add_heading_2(doc, "3.3 Hardware Requirements")
    add_p(doc, 
          "Development Environment: Modern multi-core x86_64 or ARM64 processor (e.g., Apple Silicon M-series or Intel Core i5/i7), "
          "minimum 8 GB RAM (16 GB recommended), and 2 GB available disk space for Node.js runtimes and local PostgreSQL testing.\n"
          "Production Cloud Server: Managed cloud serverless compute nodes provisioned dynamically by Vercel; zero local server hardware maintenance.")

    add_heading_2(doc, "3.4 Software Requirements")
    add_p(doc, "Table 2 details the software tools, runtimes, and platform specifications utilized across development and production.")

    tbl_sw = doc.add_table(rows=7, cols=3)
    format_table(tbl_sw, col_widths=[1.8, 2.2, 2.2])
    sw_data = [
        ("Layer / Component", "Development Specification", "Production Specification"),
        ("Operating System", "macOS Sonoma (Darwin arm64) / Linux", "Vercel Serverless Linux Container"),
        ("Runtime Environment", "Node.js 24.x (LTS) & npm 11+", "Node.js 24.x Serverless Execution"),
        ("Application Framework", "Next.js 15 (App Router) & React 19", "Next.js 15 Optimized Turbopack Bundle"),
        ("Database Engine", "PostgreSQL 16 (Local instance)", "Neon Serverless PostgreSQL 16+"),
        ("ORM & Migration Engine", "Prisma ORM 6.4.1 (Prisma Client & CLI)", "Prisma Client (Pooled PgBouncer)"),
        ("Version Control", "Git 2.40+ (Local repository)", "GitHub Canonical Source of Truth")
    ]
    for idx, (c1, c2, c3) in enumerate(sw_data):
        tbl_sw.cell(idx, 0).paragraphs[0].text = c1
        tbl_sw.cell(idx, 1).paragraphs[0].text = c2
        tbl_sw.cell(idx, 2).paragraphs[0].text = c3

    add_heading_2(doc, "3.5 User Roles & Permissions Matrix")
    add_p(doc, "Table 3 outlines the role-based access control (RBAC) boundaries enforced across the three primary system roles.")

    tbl_roles = doc.add_table(rows=4, cols=3)
    format_table(tbl_roles, col_widths=[1.2, 2.5, 2.5])
    role_data = [
        ("Role Enum", "Operational Definition", "Permitted System Actions"),
        ("USER", "Consumer or enterprise device owner seeking hardware servicing", "Register/manage owned devices, file repair tickets, run diagnostics, review bids, accept quotes, view passports, write reviews."),
        ("REPAIRER", "Verified independent technician or commercial repair workshop", "Create business profile, declare component specializations, browse open tickets, submit quotes, manage assigned jobs in workbench."),
        ("ADMIN", "Platform administrator maintaining governance and integrity", "Catalog-wide visibility, technician verification review, ticket inspection, demo seed resets, audit trail inspection.")
    ]
    for idx, (c1, c2, c3) in enumerate(role_data):
        tbl_roles.cell(idx, 0).paragraphs[0].text = c1
        tbl_roles.cell(idx, 1).paragraphs[0].text = c2
        tbl_roles.cell(idx, 2).paragraphs[0].text = c3

    add_heading_2(doc, "3.6 System Constraints & Invariants")
    add_bullet(doc, "Email Uniqueness: Every User record enforces a unique index on User.email.", bold_prefix="• Invariant 1: ")
    add_bullet(doc, "Quote Exclusivity in Jobs: A Quote can produce at most one RepairJob (enforced via unique Quote.id in RepairJob).", bold_prefix="• Invariant 2: ")
    add_bullet(doc, "Single Active Job per Request: A RepairRequest can produce exactly one active RepairJob, preventing multiple accepted bids.", bold_prefix="• Invariant 3: ")
    add_bullet(doc, "Review Uniqueness: Exactly one Review per completed RepairJob; reviews cannot be filed for uncompleted or cancelled jobs.", bold_prefix="• Invariant 4: ")
    add_bullet(doc, "Cascading Deletions: Deleting a User or Device automatically cascades to delete child tickets and quotes, preventing orphaned records.", bold_prefix="• Invariant 5: ")

    # =========================================================================
    # CHAPTER 4 — SYSTEM ARCHITECTURE
    # =========================================================================
    add_heading_1(doc, "CHAPTER 4: SYSTEM ARCHITECTURE")

    add_heading_2(doc, "4.1 Overall Architectural Topology")
    add_p(doc, 
          "RepairGraph implements a unified, full-stack cloud architecture where frontend presentation, REST API route handlers, "
          "server business services, and deterministic decision algorithms reside within a single cohesive Next.js 15 deployment "
          "on Vercel. Persistent relational storage is provided by Neon Managed Serverless PostgreSQL connected over secure pooled TLS.")
    
    add_figure_image(doc, "docs/report_assets/figure1_architecture.png", "Figure 1: RepairGraph Unified Full-Stack Architecture")

    add_heading_2(doc, "4.2 Client Layer (React 19 & Tailwind CSS v4)")
    add_p(doc, 
          "The presentation layer utilizes Next.js 15 App Router architecture with React 19. Pages are structured as Server Components "
          "by default for rapid server-side rendering (SSR) and search engine optimization, while interactive stateful interfaces "
          "(e.g., the interactive workbench stepper, symptom input modals, and quote acceptance dialogs) leverage React Client Components. "
          "Styling follows an industrial editorial design system implemented via Tailwind CSS v4, utilizing monospace tabular figures "
          "for financial accuracy, warm stone backgrounds (#fafaf9), deep charcoal typography (#1c1917), and international safety orange accents (#ea580c).")

    add_heading_2(doc, "4.3 Application & Service Layer")
    add_p(doc, 
          "The backend follows a domain-driven, layered service architecture located under src/server/services/*. Specific service modules "
          "encapsulate all business rules, lifecycle state machines, and relational transformations:")
    add_bullet(doc, "auth.service.ts: Registration, bcrypt salting, credential verification, and profile sanitization.", bold_prefix="• ")
    add_bullet(doc, "device.service.ts: Ownership verification, hardware valuation formulas, and catalog querying.", bold_prefix="• ")
    add_bullet(doc, "quote.service.ts: Bid creation, technician eligibility checks, and atomic quote acceptance workflows.", bold_prefix="• ")
    add_bullet(doc, "repairJob.service.ts: Workbench stage advancements and automated Repair Passport ledger logging.", bold_prefix="• ")
    add_bullet(doc, "repairer.service.ts: Directory queries, specialization filtering, and average star-rating recalculations.", bold_prefix="• ")

    add_heading_2(doc, "4.4 API & Routing Layer")
    add_p(doc, 
          "The REST API is implemented through 21 Next.js Route Handler files (route.ts) under src/app/api/*, serving 36 distinct HTTP method "
          "handlers. All endpoints adhere to standard REST conventions, parsing request bodies through Zod validation schemas and returning "
          "consistent JSON envelopes ({ data: ... } or RFC 7807 structured error objects).")

    add_heading_2(doc, "4.5 Authentication & Session Management")
    add_p(doc, 
          "User sessions are maintained via JSON Web Tokens (JWT) signed with HMAC-SHA256. To eliminate cross-site scripting (XSS) token "
          "theft from localStorage, tokens are transmitted exclusively via HttpOnly, SameSite=lax, Secure cookies (rg_token). Raw tokens "
          "are suppressed from JSON responses during browser login.")

    add_heading_2(doc, "4.6 Deterministic Decision Engine")
    add_p(doc, 
          "The diagnostic core resides in src/server/engine/* and operates entirely without external API calls or non-deterministic ML models. "
          "Tier 1 extracts diagnostic tokens via regex word-boundary matching, while Tier 2 evaluates the 7-Factor Repairability Score Model, "
          "depreciation math, and statutory lifecycle thresholds in sub-millisecond execution time.")

    add_heading_2(doc, "4.7 Persistence Layer (Prisma ORM)")
    add_p(doc, 
          "Database interaction is mediated by Prisma ORM 6.4.1 (@prisma/client). Prisma provides compile-time type safety, parameterized "
          "SQL generation that immunizes the application against SQL injection, and automated relational constraint enforcement across "
          "11 normalized models.")

    add_heading_2(doc, "4.8 Cloud Deployment Architecture")
    add_p(doc, 
          "The entire web application and backend API execute within Vercel Serverless Functions on the Node.js 24.x runtime. Vercel's global "
          "Anycast Edge Network handles global DNS resolution, DDoS defense, static asset caching, and TLS 1.3 termination. Persistent data "
          "resides in Neon Serverless PostgreSQL, utilizing PgBouncer connection pooling to support bursty serverless workloads without "
          "exhausting database connection limits.")

    add_heading_2(doc, "4.9 End-to-End Data & Control Flow")
    add_p(doc, 
          "Figure 2 illustrates the multi-stage validation, security, and processing pipeline traversed by every incoming client request.")
    
    add_figure_image(doc, "docs/report_assets/figure2_data_flow.png", "Figure 2: End-to-End Request & Data Processing Pipeline")

    # =========================================================================
    # CHAPTER 5 — TECHNOLOGY STACK
    # =========================================================================
    add_heading_1(doc, "CHAPTER 5: TECHNOLOGY STACK")

    add_heading_2(doc, "5.1 Framework & Runtime Selection")
    add_p(doc, 
          "Next.js 15 (App Router) was selected as the foundational full-stack framework because it uniquely unifies server-side React "
          "rendering, static asset optimization, and robust Node.js serverless route execution within a single codebase. Running on the "
          "latest Node.js 24.x LTS runtime ensures access to modern JavaScript engine optimizations (V8 v12+) and native cryptographic primitives.")

    add_heading_2(doc, "5.2 Frontend & UI Libraries")
    add_bullet(doc, "React 19: Provides concurrent rendering, transition hooks, and modular component architecture.", bold_prefix="• ")
    add_bullet(doc, "Tailwind CSS v4: Delivers an ultra-fast post-CSS engine and utility-first styling without runtime overhead.", bold_prefix="• ")
    add_bullet(doc, "Lucide React (v1.45.0): Lightweight, accessible SVG iconography representing hardware components and status states.", bold_prefix="• ")

    add_heading_2(doc, "5.3 Backend & Validation Technologies")
    add_bullet(doc, "TypeScript 5: Enforces static typing across the entire codebase, eliminating null pointer exceptions and schema mismatch bugs.", bold_prefix="• ")
    add_bullet(doc, "Zod (v4.6.2): Type-safe schema validation library used across all route handlers to strictly parse and sanitize incoming HTTP payloads.", bold_prefix="• ")
    add_bullet(doc, "bcryptjs (v3.0.3): Implements salted Blowfish hashing (10 rounds) for cryptographically secure user password storage.", bold_prefix="• ")
    add_bullet(doc, "jsonwebtoken (v9.0.3): Implements standard RFC 7519 JSON Web Token signing and verification with HMAC-SHA256.", bold_prefix="• ")

    add_heading_2(doc, "5.4 Database & Persistence Stack")
    add_bullet(doc, "PostgreSQL 16+: Advanced open-source object-relational database providing ACID guarantees, complex join capabilities, and JSONB support.", bold_prefix="• ")
    add_bullet(doc, "Prisma ORM (v6.4.1): Modern type-safe database toolkit providing declarative schema modeling, migration automation, and query building.", bold_prefix="• ")

    add_heading_2(doc, "5.5 Cloud Infrastructure Providers")
    add_bullet(doc, "Vercel Serverless Platform: Cloud hosting platform providing automated Git integration, serverless function scaling, and Anycast CDN caching.", bold_prefix="• ")
    add_bullet(doc, "Neon Serverless PostgreSQL: Fully managed cloud PostgreSQL offering automated connection pooling (PgBouncer) and storage auto-scaling.", bold_prefix="• ")

    add_heading_2(doc, "5.6 Technology Stack Inventory Table")
    add_p(doc, "Table 4 provides a comprehensive inventory of the technologies, versions, and architectural roles across RepairGraph.")

    tbl_tech = doc.add_table(rows=11, cols=4)
    format_table(tbl_tech, col_widths=[1.5, 1.0, 1.5, 2.2])
    tech_data = [
        ("Technology / Tool", "Version", "Category", "Architectural Role"),
        ("Next.js", "15.5.25", "Full-Stack Framework", "App Router frontend, SSR/SSG rendering & API Route Handlers"),
        ("React", "19.1.0", "UI Library", "Component-driven presentation & client-side hydration"),
        ("TypeScript", "^5.0.0", "Programming Language", "Static typing across data models, services, and UI components"),
        ("Node.js", "24.x (LTS)", "Runtime Environment", "Serverless function execution on Vercel runtime"),
        ("Tailwind CSS", "^4.0.0", "CSS Framework", "Industrial editorial design system and responsive layout styling"),
        ("Prisma ORM", "^6.4.1", "Data Access & ORM", "Type-safe database client, schema migrations, and SQL generation"),
        ("PostgreSQL (Neon)", "16+", "Relational Database", "Cloud-hosted persistent storage with ACID invariants and pooling"),
        ("Vercel", "Production", "Cloud Platform", "Continuous deployment, serverless compute, and Anycast CDN"),
        ("Zod", "^4.6.2", "Schema Validation", "Input validation and runtime schema boundary enforcement"),
        ("bcryptjs & JWT", "^3.0.3 / ^9.0.3", "Security & Auth", "Password salting (10 rounds) & HMAC-SHA256 session tokens")
    ]
    for idx, (c1, c2, c3, c4) in enumerate(tech_data):
        tbl_tech.cell(idx, 0).paragraphs[0].text = c1
        tbl_tech.cell(idx, 1).paragraphs[0].text = c2
        tbl_tech.cell(idx, 2).paragraphs[0].text = c3
        tbl_tech.cell(idx, 3).paragraphs[0].text = c4

    doc.add_page_break()
